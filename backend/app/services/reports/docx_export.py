import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.schemas.report import ReportContent

ACCENT_COLOR = RGBColor(0x1F, 0x4E, 0x79)  # dark blue, used for title + headings
HEADER_FILL = "1F4E79"  # same color, hex without '#', for table header shading

_INLINE_PATTERN = re.compile(
    r'(\*\*.+?\*\*|\*.+?\*|<mark color="\w+">.*?</mark>|<color name="\w+">.*?</color>)', re.DOTALL
)
_MARK_PATTERN = re.compile(r'<mark color="(\w+)">(.*?)</mark>', re.DOTALL)
_COLOR_PATTERN = re.compile(r'<color name="(\w+)">(.*?)</color>', re.DOTALL)

_HIGHLIGHT_COLORS = {
    "yellow": WD_COLOR_INDEX.YELLOW,
    "green": WD_COLOR_INDEX.BRIGHT_GREEN,
    "red": WD_COLOR_INDEX.RED,
    "blue": WD_COLOR_INDEX.BLUE,
    "pink": WD_COLOR_INDEX.PINK,
    "gray": WD_COLOR_INDEX.GRAY_25,
}

# Font colors are arbitrary RGB, unlike highlights which are limited to
# Word's fixed highlighter palette — so this map isn't constrained to the
# same color set as _HIGHLIGHT_COLORS.
_FONT_COLORS = {
    "yellow": RGBColor(0xC9, 0xA2, 0x00),
    "green": RGBColor(0x1E, 0x7B, 0x34),
    "red": RGBColor(0xC0, 0x1C, 0x1C),
    "blue": RGBColor(0x1F, 0x4E, 0x79),
    "pink": RGBColor(0xC2, 0x18, 0x5B),
    "gray": RGBColor(0x59, 0x59, 0x59),
    "orange": RGBColor(0xD9, 0x73, 0x0D),
    "purple": RGBColor(0x6A, 0x1B, 0x9A),
    "black": RGBColor(0x00, 0x00, 0x00),
}


def _add_page_number(section) -> None:
    """Insert a live page-number field into the section's footer."""
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def _shade_cell(cell, hex_color: str) -> None:
    """python-docx has no high-level API for cell shading, so we build the
    raw <w:shd> XML element and attach it to the cell's properties."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def _render_marks_into_paragraph(paragraph, text: str) -> None:
    """Parses **bold**, *italic*, <mark>, and <color> spans in `text` and
    adds each piece as its own run with the right formatting applied.
    Works on any paragraph — a body paragraph, a table cell's paragraph,
    or a heading — since all of these are the same python-docx type."""
    pos = 0
    for match in _INLINE_PATTERN.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])

        token = match.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("<mark"):
            m = _MARK_PATTERN.match(token)
            run = paragraph.add_run(m.group(2))
            run.font.highlight_color = _HIGHLIGHT_COLORS.get(m.group(1).lower(), WD_COLOR_INDEX.YELLOW)
        elif token.startswith("<color"):
            m = _COLOR_PATTERN.match(token)
            run = paragraph.add_run(m.group(2))
            run.font.color.rgb = _FONT_COLORS.get(m.group(1).lower(), RGBColor(0x00, 0x00, 0x00))
        else:
            paragraph.add_run(token[1:-1]).italic = True

        pos = match.end()

    if pos < len(text):
        paragraph.add_run(text[pos:])


def _add_paragraph_with_marks(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    _render_marks_into_paragraph(paragraph, text)
    return paragraph


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        para = header_cells[i].paragraphs[0]
        _render_marks_into_paragraph(para, header_text)
        _shade_cell(header_cells[i], HEADER_FILL)
        for run in para.runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.bold = True

    for row_data in rows:
        row_cells = table.add_row().cells
        for i, value in enumerate(row_data):
            _render_marks_into_paragraph(row_cells[i].paragraphs[0], value)

    doc.add_paragraph()
def _add_images(doc: Document, image_paths: list[Path]) -> None:
    """Lay out a section's images to avoid wasting pages: a single image
    gets a moderate width; multiple images share rows via a borderless
    table, with column count adapting to how many there are."""
    if not image_paths:
        return

    if len(image_paths) == 1:
        doc.add_picture(str(image_paths[0]), width=Inches(4))
        return

    columns = 2 if len(image_paths) <= 4 else 3
    cell_width = Inches(6 / columns)

    rows_needed = (len(image_paths) + columns - 1) // columns
    grid = doc.add_table(rows=rows_needed, cols=columns)
    grid.autofit = True

    for idx, img_path in enumerate(image_paths):
        cell = grid.rows[idx // columns].cells[idx % columns]
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run()
        run.add_picture(str(img_path), width=cell_width)

    doc.add_paragraph()


def build_docx(content: ReportContent, image_paths_by_id: dict[str, Path], output_path: Path) -> Path:
    doc = Document()

    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    _add_page_number(section)

    title_paragraph = doc.add_heading(content.title, level=0)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title_paragraph.runs:
        run.font.color.rgb = ACCENT_COLOR

    for sec in content.sections:
        heading_paragraph = doc.add_heading("", level=1)
        _render_marks_into_paragraph(heading_paragraph, sec.heading)
        for run in heading_paragraph.runs:
            if not run.font.color.rgb:
                run.font.color.rgb = ACCENT_COLOR

        for para_text in sec.body.split("\n\n"):
            if para_text.strip():
                _add_paragraph_with_marks(doc, para_text.strip())

        if sec.table:
            _add_table(doc, sec.table.headers, sec.table.rows)

        section_image_paths = [
            image_paths_by_id[fid] for fid in sec.image_file_ids if fid in image_paths_by_id
        ]
        _add_images(doc, section_image_paths)

    doc.save(output_path)
    return output_path
