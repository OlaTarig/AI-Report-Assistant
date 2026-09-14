from docx import Document


def extract(path) -> dict:
    doc = Document(path)

    lines = []
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        # Paragraph style names like "Heading 1", "Heading 2", "Title" tell
        # us this line is structural, not body text — mark it explicitly so
        # the AI can tell the difference in the flattened text we send it.
        if p.style and p.style.name and p.style.name.startswith(("Heading", "Title")):
            lines.append(f"[{p.style.name.upper()}] {p.text}")
        else:
            lines.append(p.text)

    tables_text = []
    for table in doc.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        tables_text.append("[TABLE]\n" + "\n".join(rows))

    text_summary = "\n".join(lines)
    if tables_text:
        text_summary += "\n\n" + "\n\n".join(tables_text)

    return {
        "paragraph_count": len(lines),
        "table_count": len(doc.tables),
        "text_summary": text_summary.strip(),
    }
